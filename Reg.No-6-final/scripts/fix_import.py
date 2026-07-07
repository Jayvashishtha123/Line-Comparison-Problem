#!/usr/bin/env python3
"""
Regenerate portfolio_dummy.json from Excel with proper value extraction (no formulae).
"""

import json
from pathlib import Path
from datetime import datetime

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    exit(1)

DATA_DIR = Path(__file__).parent.parent / "data"
EXCEL_FILE = DATA_DIR / "Worksheet PMO App Hackathon  1 1.xlsx"

def get_cell_value(cell):
    """Extract actual value from cell, not formula."""
    if cell is None:
        return None
    
    value = cell.value
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    
    # Skip formula strings
    if isinstance(value, str) and value.startswith('='):
        return None
    
    # Convert to string and strip
    return str(value).strip() if value else None


def extract_projects():
    """Extract projects from Excel sheet."""
    if not EXCEL_FILE.exists():
        print(f"ERROR: Excel file not found: {EXCEL_FILE}")
        exit(1)
    
    print(f"Loading Excel file: {EXCEL_FILE}")
    # Use data_only=True to get computed formula values instead of formulas
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    
    # Find the "Program 2026" sheet
    ws = wb["Program 2026 "]
    
    # Column mapping (0-indexed from Excel)
    COLUMN_MAP = {
        "project_status": 0,              # Col 1: Project Status
        "initiative_name": 1,             # Col 2: Initiative Name
        "project_name": 2,                # Col 3: Project Name
        "description": 3,                 # Col 4: Description
        "dt_team_accountable": 4,         # Col 5: DT Team Accountable
        "project_otd_status": 5,          # Col 6: OTD Status
        "state": 6,                       # Col 7: State
        "project_otd_phase": 7,           # Col 8: Phase
        "estd_start_date": 8,             # Col 9: Estd. Start Date
        "estd_end_date": 9,               # Col 10: Estd. End Date
        "adjusted_end_date": 10,          # Col 11: Adjusted End Date
        "category": 11,                   # Col 12: Category
        "business_unit": 12,              # Col 13: Business Unit
        "primary_function": 13,           # Col 14: Primary Function
        "region": 14,                     # Col 15: Region
        "executive_sponsor": 15,          # Col 16: Executive Sponsor
        "budget_2026_flag": 16,           # Col 17: 2026 budget?
        "benefits_plan_2026": 17,         # Col 18: 2026 Benefits Plan
        "dt_teams_involved": 19,          # Col 20: DT Teams Involved
        "filter_key": 20,                 # Col 21: For Filter Purposes
        "dt_lead": 21,                    # Col 22: DT lead
        "business_project_lead": 22,      # Col 23: Business Project Lead
        "level_of_effort": 23,            # Col 24: Level of Effort
        "technology": 24,                 # Col 25: Technology
        "budget": 25,                     # Col 26: Budget
    }
    
    projects = []
    row_num = 2
    
    while row_num <= 300:
        # Keep rows if any important leading columns are present
        row_key_values = [
            get_cell_value(ws.cell(row=row_num, column=2)),  # initiative
            get_cell_value(ws.cell(row_num, column=3)),  # project name
            get_cell_value(ws.cell(row_num, column=5)),  # accountable team
            get_cell_value(ws.cell(row_num, column=7)),  # state
        ]

        if not any(v for v in row_key_values):
            row_num += 1
            continue
        
        # Build project dict
        project = {
            "project_id": f"PRJ-{len(projects) + 1:04d}"
        }
        
        # Extract each field
        for field_name, col_idx in COLUMN_MAP.items():
            col_num = col_idx + 1
            cell = ws.cell(row=row_num, column=col_num)
            value = get_cell_value(cell)
            project[field_name] = value
        
        # Only add if has project name
        if project.get('project_name'):
            projects.append(project)
            print(f"Row {row_num}: {project.get('project_name')[:50]}")
        
        row_num += 1
    
    print(f"\nSuccessfully extracted {len(projects)} projects")
    return projects


def main():
    projects = extract_projects()
    
    # Prepare output
    output = {
        "metadata": {
            "source_workbook": "Worksheet PMO App Hackathon  1 1.xlsx",
            "source_sheet": "Program 2026 ",
            "description": "Complete development dataset from Program 2026 sheet",
            "generated_on": datetime.now().strftime("%Y-%m-%d"),
            "record_count": len(projects)
        },
        "projects": projects
    }
    
    # Write to portfolio_dummy.json
    output_file = DATA_DIR / "portfolio_dummy.json"
    output_file.write_text(json.dumps(output, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"\nWrote {output_file}")
    
    # Also delete portfolio.json to force reload
    working_file = DATA_DIR / "portfolio.json"
    if working_file.exists():
        working_file.unlink()
        print(f"Deleted {working_file} (will be recreated)")
    
    print("\nDone! Start the app to reload data.")


if __name__ == "__main__":
    main()
