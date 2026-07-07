import openpyxl
wb = openpyxl.load_workbook('data/Worksheet PMO App Hackathon  1 1.xlsx')
ws = wb['Program 2026 ']

print('Headers (Row 1):')
for i in range(1, 26):
    cell = ws.cell(row=1, column=i)
    print(f'  Col {i:2d}: {cell.value}')

print('\nRow 2 values:')
for i in range(1, 26):
    cell = ws.cell(row=2, column=i)
    print(f'  Col {i:2d}: {cell.value}')
