"""Business logic service for portfolio projects."""

from __future__ import annotations

import json
import os
import re
import uuid
import base64
import csv
from difflib import SequenceMatcher
from io import BytesIO, StringIO
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import from_excel

from services.audit_trail_service import AuditTrailService
from services.validation_service import ProjectFormValidator


class ProjectService:
    """Service-layer operations for project CRUD, search, grouping, and status workflows."""

    SEARCH_FIELDS = [
        "initiative_name",
        "project_name",
        "description",
        "dt_team_accountable",
        "state",
        "project_otd_phase",
        "category",
        "region",
        "executive_sponsor",
    ]
    BLANK_FILTER_TOKEN = "__BLANK__"

    @staticmethod
    def _normalize_header_key(value: str) -> str:
        text = str(value or "").strip().lower()
        return re.sub(r"[^a-z0-9]", "", text)

    def get_active_budget_years(self) -> Tuple[int, int]:
        """Return budget/benefits years shown in UI: current and next year."""
        current_year = datetime.now().year
        return current_year, current_year + 1

    def get_budget_field_names_for_ui(self) -> Dict[str, str]:
        """Return dynamic budget field keys for current and next UI years."""
        current_year, next_year = self.get_active_budget_years()
        return {
            "current_budget": f"budget_{current_year}_flag",
            "current_benefits": f"benefits_plan_{current_year}",
            "next_budget": f"budget_{next_year}_flag",
            "next_benefits": f"benefits_plan_{next_year}",
        }

    def __init__(self, data_dir: str = "data"):
        """
        Initialize project service with data directory.
        
        Args:
            data_dir: Path to data directory containing JSON files
        """
        self.data_dir = Path(data_dir)
        self.portfolio_file = self.data_dir / "portfolio.json"
        self.form_spec_file = self.data_dir / "project_form_spec.json"
        self.portfolio_dummy_file = self.data_dir / "portfolio_dummy.json"
        self.notification_log_file = self.data_dir / "notification_log.json"

        self.form_spec = self._load_json(self.form_spec_file)
        self.validator = ProjectFormValidator(self.form_spec)
        self.audit_service = AuditTrailService(data_dir)
        self.portfolio = self._load_or_bootstrap_portfolio()

    def _load_json(self, filepath: Path) -> Dict[str, Any]:
        if not filepath.exists():
            return {}
        try:
            return json.loads(filepath.read_text(encoding="utf-8"))
        except Exception:
            return {}

    def _write_json(self, filepath: Path, data: Any) -> bool:
        try:
            filepath.parent.mkdir(parents=True, exist_ok=True)
            filepath.write_text(json.dumps(data, indent=2), encoding="utf-8")
            return True
        except Exception:
            return False

    def _load_or_bootstrap_portfolio(self) -> Dict[str, Any]:
        now = datetime.now().isoformat()
        if self.portfolio_file.exists():
            data = self._load_json(self.portfolio_file)
            if isinstance(data, dict) and "projects" in data:
                if "metadata" not in data:
                    data["metadata"] = {"created_at": now, "last_updated": now}
                return data

        projects: List[Dict[str, Any]] = []
        dummy = self._load_json(self.portfolio_dummy_file)
        if isinstance(dummy, dict):
            projects = list(dummy.get("projects", []))

        seeded = {
            "projects": projects,
            "metadata": {
                "created_at": now,
                "last_updated": now,
                "seeded_from_dummy": True,
            },
        }
        self._write_json(self.portfolio_file, seeded)
        return seeded

    def _save_portfolio(self) -> bool:
        self.portfolio.setdefault("metadata", {})
        self.portfolio["metadata"]["last_updated"] = datetime.now().isoformat()
        return self._write_json(self.portfolio_file, self.portfolio)

    def _norm(self, value: Any) -> str:
        return " ".join(str(value or "").lower().split())

    def _is_duplicate(
        self,
        project_name: str,
        initiative_name: str,
        ignore_project_id: Optional[str] = None,
    ) -> Optional[Dict[str, Any]]:
        pn = self._norm(project_name)
        inn = self._norm(initiative_name)
        for p in self.portfolio.get("projects", []):
            if ignore_project_id and p.get("project_id") == ignore_project_id:
                continue
            if self._norm(p.get("project_name")) == pn and self._norm(p.get("initiative_name")) == inn:
                return p
        return None

    def validate_name_pair(self, project_name: str, initiative_name: str) -> Tuple[bool, Optional[str]]:
        """Validate project/initiative name pair uniqueness for add-form live checks."""
        project_name = (project_name or "").strip()
        initiative_name = (initiative_name or "").strip()
        if not project_name or not initiative_name:
            return True, None

        dup = self._is_duplicate(project_name, initiative_name)
        if not dup:
            return True, None

        return False, f"Duplicate project exists: {dup.get('project_id')} - {dup.get('project_name')}"

    def validate_project_form(self, form_data: Dict[str, Any]) -> Tuple[bool, List[Dict[str, str]]]:
        """
        Validate project form submission.
        
        Args:
            form_data: Form data dictionary
            
        Returns:
            Tuple of (is_valid, list of error dicts with 'field' and 'message' keys)
        """
        is_valid, errors = self.validator.validate_form(form_data)
        
        return is_valid, [{"field": error.field, "message": error.message} for error in errors]

    def create_project(self, form_data: Dict[str, Any]) -> Tuple[bool, Optional[Dict[str, Any]], List[Dict[str, str]]]:
        """
        Create a new project with validation and defaults.
        
        Args:
            form_data: Raw form data from user
            
        Returns:
            Tuple of (success, project_record or None, list of errors)
        """
        # Apply defaults before validation
        project_data = self.validator.set_defaults(form_data.copy())
        project_data = self.validator.set_system_generated_values(project_data)

        dup = self._is_duplicate(project_data.get("project_name"), project_data.get("initiative_name"))
        if dup:
            return False, None, [{
                "field": "project_name",
                "message": f"Duplicate project exists: {dup.get('project_id')} - {dup.get('project_name')}"
            }]

        is_valid, errors = self.validate_project_form(project_data)
        if not is_valid:
            return False, None, errors
        
        # Generate project ID
        project_id = f"PRJ-{uuid.uuid4().hex[:8].upper()}"
        project_data["project_id"] = project_id
        
        # Set creation timestamp
        now = datetime.now().isoformat()
        project_data["created_at"] = now
        project_data["updated_at"] = now
        
        self.portfolio.setdefault("projects", []).append(project_data)
        
        # Persist to disk
        if not self._save_portfolio():
            return False, None, [{"field": "general", "message": "Failed to save project"}]
        
        # Log project creation in audit trail
        self.audit_service.log_project_creation(
            project_id, 
            project_data, 
            user_id=os.getenv("CURRENT_USER", "System"),
            user_name=os.getenv("CURRENT_USER_NAME", "System")
        )
        
        return True, project_data, []

    def _importable_field_ids(self) -> List[str]:
        field_ids: List[str] = []
        sections = self.form_spec.get("project_form", {}).get("sections", [])
        for section in sections:
            for field in section.get("fields", []):
                if field.get("is_system_generated"):
                    continue
                field_id = str(field.get("field_id") or "").strip()
                if field_id and field_id not in field_ids:
                    field_ids.append(field_id)

        for dynamic_field in self.get_budget_field_names_for_ui().values():
            if dynamic_field not in field_ids:
                field_ids.append(dynamic_field)

        return field_ids

    def _import_header_alias_map(self) -> Dict[str, str]:
        alias_map: Dict[str, str] = {}
        sections = self.form_spec.get("project_form", {}).get("sections", [])
        for section in sections:
            for field in section.get("fields", []):
                if field.get("is_system_generated"):
                    continue

                field_id = str(field.get("field_id") or "").strip()
                label = str(field.get("label") or "").strip()
                if not field_id:
                    continue

                alias_map[self._normalize_header_key(field_id)] = field_id
                if label:
                    alias_map[self._normalize_header_key(label)] = field_id

        for dynamic_field in self.get_budget_field_names_for_ui().values():
            alias_map[self._normalize_header_key(dynamic_field)] = dynamic_field

        return alias_map

    def _date_field_ids(self) -> set:
        result = set()
        sections = self.form_spec.get("project_form", {}).get("sections", [])
        for section in sections:
            for field in section.get("fields", []):
                if field.get("type") == "date":
                    field_id = str(field.get("field_id") or "").strip()
                    if field_id:
                        result.add(field_id)
        return result

    def _normalize_excel_cell(self, value: Any, is_date_field: bool, workbook_epoch: datetime = None) -> str:
        if value is None:
            return ""

        if is_date_field:
            parsed_date = None
            if isinstance(value, datetime):
                parsed_date = value.date()
            elif isinstance(value, date):
                parsed_date = value
            elif isinstance(value, (int, float)):
                try:
                    parsed_date = from_excel(value, epoch=workbook_epoch).date() if workbook_epoch else from_excel(value).date()
                except Exception:
                    parsed_date = None
            else:
                text = str(value).strip()
                if not text:
                    return ""
                for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y", "%d/%m/%Y", "%Y/%m/%d"):
                    try:
                        parsed_date = datetime.strptime(text, fmt).date()
                        break
                    except Exception:
                        continue

            if parsed_date is not None:
                return parsed_date.strftime("%Y-%m-%d")
            return str(value).strip()

        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d")
        if isinstance(value, date):
            return value.strftime("%Y-%m-%d")
        return str(value).strip()

    def import_projects_from_excel(self, file_stream: Any, user_id: str = "System", user_name: str = "System") -> Dict[str, Any]:
        try:
            workbook = load_workbook(filename=file_stream, data_only=True)
        except Exception:
            return {
                "success": False,
                "error": "Unable to read Excel file. Please upload a valid .xlsx file.",
            }

        worksheet = workbook.active
        if worksheet.max_row < 2:
            return {
                "success": False,
                "error": "No data rows found in the uploaded file.",
            }

        field_ids = self._importable_field_ids()
        date_fields = self._date_field_ids()

        header_cells = list(worksheet[1])
        original_headers = [str(cell.value) if cell.value is not None else "" for cell in header_cells]
        header_to_field: Dict[int, str] = {}
        normalized_field_map = self._import_header_alias_map()
        error_details_header_key = self._normalize_header_key("error details")
        existing_error_col_idx = next(
            (
                idx
                for idx, header in enumerate(original_headers)
                if self._normalize_header_key(header) == error_details_header_key
            ),
            -1,
        )

        for idx, cell in enumerate(header_cells):
            header_value = str(cell.value or "").strip()
            if not header_value:
                continue
            # Skip re-import metadata column from failed export files.
            if self._normalize_header_key(header_value) == error_details_header_key:
                continue
            mapped = normalized_field_map.get(self._normalize_header_key(header_value))
            if mapped:
                header_to_field[idx] = mapped

        if not header_to_field:
            return {
                "success": False,
                "error": "No matching columns found. Ensure headers match project field names.",
            }

        total_rows = 0
        imported_count = 0
        failed_rows: List[Dict[str, str]] = []
        failed_export_rows: List[List[str]] = []

        previous_user = os.getenv("CURRENT_USER")
        previous_user_name = os.getenv("CURRENT_USER_NAME")
        os.environ["CURRENT_USER"] = user_id or "System"
        os.environ["CURRENT_USER_NAME"] = user_name or user_id or "System"

        try:
            for row_number, row_cells in enumerate(worksheet.iter_rows(min_row=2), start=2):
                values_present = False
                row_data: Dict[str, str] = {field_id: "" for field_id in field_ids}
                original_row_values = [
                    self._normalize_excel_cell(
                        value=cell.value,
                        is_date_field=False,
                        workbook_epoch=getattr(workbook, "epoch", None),
                    )
                    for cell in row_cells[:len(original_headers)]
                ]

                if len(original_row_values) < len(original_headers):
                    original_row_values.extend([""] * (len(original_headers) - len(original_row_values)))

                for col_idx, cell in enumerate(row_cells):
                    mapped_field = header_to_field.get(col_idx)
                    if not mapped_field:
                        continue
                    normalized = self._normalize_excel_cell(
                        value=cell.value,
                        is_date_field=mapped_field in date_fields,
                        workbook_epoch=getattr(workbook, "epoch", None),
                    )
                    if normalized != "":
                        values_present = True
                    row_data[mapped_field] = normalized

                if not values_present:
                    continue

                total_rows += 1
                success, project, errors = self.create_project(row_data)
                if success:
                    imported_count += 1
                    continue

                project_name = (row_data.get("project_name") or "").strip()
                reason = " | ".join(f"{e.get('field')}: {e.get('message')}" for e in errors) if errors else "Unknown error"
                failed_rows.append({
                    "row_number": str(row_number),
                    "project_name": project_name,
                    "reason": reason,
                })
                if existing_error_col_idx >= 0:
                    while len(original_row_values) <= existing_error_col_idx:
                        original_row_values.append("")
                    original_row_values[existing_error_col_idx] = reason
                    failed_export_rows.append(original_row_values)
                else:
                    failed_export_rows.append(original_row_values + [reason])
        finally:
            if previous_user is None:
                os.environ.pop("CURRENT_USER", None)
            else:
                os.environ["CURRENT_USER"] = previous_user

            if previous_user_name is None:
                os.environ.pop("CURRENT_USER_NAME", None)
            else:
                os.environ["CURRENT_USER_NAME"] = previous_user_name

        failed_count = len(failed_rows)

        failed_export_file_b64 = ""
        failed_export_filename = ""
        if failed_count > 0:
            export_wb = Workbook()
            export_ws = export_wb.active
            export_ws.title = "Failed Rows"
            if existing_error_col_idx >= 0:
                export_ws.append(original_headers)
            else:
                export_ws.append(original_headers + ["Error Details"])
            for row in failed_export_rows:
                export_ws.append(row)

            stream = BytesIO()
            export_wb.save(stream)
            failed_export_file_b64 = base64.b64encode(stream.getvalue()).decode("ascii")
            failed_export_filename = f"failed_import_rows_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

        return {
            "success": True,
            "total_rows": total_rows,
            "imported_count": imported_count,
            "failed_count": failed_count,
            "message": f"Imported {imported_count} of {total_rows}. {failed_count} rows failed",
            "failed_rows": failed_rows,
            "failed_export_file_base64": failed_export_file_b64,
            "failed_export_file_name": failed_export_filename,
        }

    def import_projects_from_file(self, file_stream: Any, filename: str, user_id: str = "System", user_name: str = "System") -> Dict[str, Any]:
        lower_name = str(filename or "").strip().lower()

        if lower_name.endswith(".xlsx"):
            try:
                file_stream.seek(0)
            except Exception:
                pass
            return self.import_projects_from_excel(file_stream, user_id=user_id, user_name=user_name)

        if lower_name.endswith(".csv"):
            try:
                file_stream.seek(0)
            except Exception:
                pass

            raw = file_stream.read()
            if raw is None:
                raw = b""

            text = ""
            if isinstance(raw, str):
                text = raw
            else:
                for enc in ("utf-8-sig", "utf-8", "latin-1"):
                    try:
                        text = raw.decode(enc)
                        break
                    except Exception:
                        continue

            rows = list(csv.reader(StringIO(text)))
            if not rows or len(rows) < 2:
                return {
                    "success": False,
                    "error": "No data rows found in the uploaded file.",
                }

            temp_wb = Workbook()
            temp_ws = temp_wb.active
            for row in rows:
                temp_ws.append(row)

            temp_stream = BytesIO()
            temp_wb.save(temp_stream)
            temp_stream.seek(0)
            return self.import_projects_from_excel(temp_stream, user_id=user_id, user_name=user_name)

        return {
            "success": False,
            "error": "Only .xlsx and .csv files are supported.",
        }

    def get_project(self, project_id: str) -> Optional[Dict[str, Any]]:
        """
        Retrieve a project by ID.
        
        Args:
            project_id: Project ID
            
        Returns:
            Project record or None if not found
        """
        for project in self.portfolio.get("projects", []):
            if project.get("project_id") == project_id:
                return project
        
        return None

    def update_project(self, project_id: str, form_data: Dict[str, Any]) -> Tuple[bool, Optional[Dict[str, Any]], List[Dict[str, str]]]:
        """
        Update an existing project.
        
        Args:
            project_id: Project ID to update
            form_data: Updated form data
            
        Returns:
            Tuple of (success, updated_project or None, list of errors)
        """
        project = self.get_project(project_id)
        if not project:
            return False, None, [{"field": "general", "message": f"Project {project_id} not found"}]

        updated_project = project.copy()
        updated_project.update(form_data)
        updated_project = self.validator.set_defaults(updated_project)
        updated_project = self.validator.set_system_generated_values(updated_project)

        dup = self._is_duplicate(
            updated_project.get("project_name"),
            updated_project.get("initiative_name"),
            ignore_project_id=project_id,
        )
        if dup:
            return False, None, [{"field": "project_name", "message": "Duplicate project exists"}]

        is_valid, errors = self.validate_project_form(updated_project)
        if not is_valid:
            return False, None, errors

        updated_project["project_id"] = project_id
        updated_project["created_at"] = project.get("created_at", datetime.now().isoformat())
        updated_project["updated_at"] = datetime.now().isoformat()

        dt_team = updated_project.get("dt_team_accountable", "")
        primary_func = updated_project.get("primary_function", "")
        updated_project["filter_key"] = f"{dt_team}{primary_func}".strip()

        for i, proj in enumerate(self.portfolio.get("projects", [])):
            if proj.get("project_id") == project_id:
                self.portfolio["projects"][i] = updated_project
                break

        if not self._save_portfolio():
            return False, None, [{"field": "general", "message": "Failed to save project"}]

        for field_key, new_value in updated_project.items():
            old_value = project.get(field_key)
            if old_value != new_value:
                self.audit_service.log_field_change(
                    project_id,
                    field_key,
                    old_value,
                    new_value,
                    user_id=os.getenv("CURRENT_USER", "Unknown"),
                    user_name=os.getenv("CURRENT_USER_NAME", os.getenv("CURRENT_USER", "Unknown")),
                    change_type="update"
                )
        
        return True, updated_project, []

    def delete_project(self, project_id: str) -> Tuple[bool, str]:
        """
        Delete a project by ID.
        
        Args:
            project_id: Project ID to delete
            
        Returns:
            Tuple of (success, message)
        """
        original_count = len(self.portfolio.get("projects", []))
        self.portfolio["projects"] = [
            proj for proj in self.portfolio.get("projects", [])
            if proj.get("project_id") != project_id
        ]
        if len(self.portfolio["projects"]) == original_count:
            return False, f"Project {project_id} not found"
        if not self._save_portfolio():
            return False, "Failed to save portfolio after deletion"
        self.audit_service.log_field_change(project_id, "project", "active", "deleted", "System", "System", "delete")
        return True, f"Project {project_id} deleted successfully"

    def get_all_projects(self) -> List[Dict[str, Any]]:
        """Get all projects."""
        return self.portfolio.get("projects", [])

    def get_form_lovs(self) -> Dict[str, List[str]]:
        """
        Get Lists of Values for form dropdowns.
        
        Returns:
            Dictionary of LOVs keyed by field name
        """
        return self.form_spec.get("form_lovs", {})

    def get_form_spec(self) -> Dict[str, Any]:
        """
        Get complete form specification.
        
        Returns:
            Form specification dictionary
        """
        return self.form_spec

    def _normalize_otd(self, otd: Optional[str]) -> Optional[str]:
        """Normalize OTD status value to canonical form matching UI options."""
        if not otd:
            return otd
        mapping = {
            "on track": "On Track",
            "on-track": "On Track",
            "adjusted plan": "Adjusted plan",
            "adjusted_plan": "Adjusted plan",
            "slip": "Slip",
        }
        return mapping.get(str(otd).strip().lower(), otd)

    def _normalize_project_status(self, status: Optional[str]) -> str:
        """Normalize project status to canonical values: G, R, Y, NA, Hold."""
        if status is None:
            return "NA"
        s = str(status).strip()
        if not s:
            return "NA"
        low = s.lower()
        mapping = {
            "g": "G",
            "green": "G",
            "r": "R",
            "red": "R",
            "a": "Y",
            "amber": "Y",
            "y": "Y",
            "yellow": "Y",
            "na": "NA",
            "n/a": "NA",
            "hold": "Hold",
        }
        return mapping.get(low, s)

    def _display_project_status(self, project: Dict[str, Any]) -> str:
        # Always display status; show NA when missing/blank.
        return self._normalize_project_status(project.get("project_status"))

    def _display_otd_status(self, project: Dict[str, Any]) -> Optional[str]:
        if project.get("state") in ["Funnel", "Evaluation", "Hold"]:
            return None
        return project.get("project_otd_status")

    def get_status_dialog_data(self, project_id: str) -> Tuple[bool, Optional[Dict[str, Any]], str]:
        project = self.get_project(project_id)
        if not project:
            return False, None, "Project not found"

        state = project.get("state") or "Evaluation"
        payload = {
            "project_id": project_id,
            "initiative_name": project.get("initiative_name", ""),
            "project_name": project.get("project_name", ""),
            "state": state,
            "project_status": self._normalize_project_status(project.get("project_status")),
            "project_otd_status": self._normalize_otd(self._display_otd_status(project)),
            "project_otd_phase": project.get("project_otd_phase") or "Discovery",
            "adjusted_end_date": project.get("adjusted_end_date", ""),
            "comments": project.get("status_comments", ""),
            "rules": {
                "project_status_blank_when_state_in": [],
                "otd_status_null_when_state_in": ["Funnel", "Evaluation", "Hold"],
                "phase_optional_when_state_in": ["Funnel"],
                "adjusted_end_date_visible_when_otd_status": "Adjusted plan",
                "adjusted_end_date_required_when_otd_status": "Adjusted plan",
            },
        }
        return True, payload, "OK"

    def update_project_status(self, project_id: str, payload: Dict[str, Any], user_id: str = None) -> Tuple[bool, Optional[Dict[str, Any]], str]:
        """
        Update project status via the Update button with audit logging.
        
        Args:
            project_id: Project ID
            payload: Status-dialog payload
            user_id: User making the change
            
        Returns:
            Tuple of (success, updated_project, message)
        """
        project = self.get_project(project_id)
        if not project:
            return False, None, f"Project {project_id} not found"

        state = payload.get("state", project.get("state"))
        # Accept whatever status the user explicitly provides; fall back to existing
        incoming_status = self._normalize_project_status(
            payload.get("project_status") if "project_status" in payload else project.get("project_status", "")
        )
        incoming_otd = self._normalize_otd(payload.get("project_otd_status") if "project_otd_status" in payload else project.get("project_otd_status"))
        incoming_phase = payload.get("project_otd_phase") if "project_otd_phase" in payload else project.get("project_otd_phase")
        adjusted = payload.get("adjusted_end_date", project.get("adjusted_end_date", ""))
        incoming_comments = payload.get("comments") if "comments" in payload else project.get("status_comments", "")

        # Phase defaults to Discovery when not provided for non-Funnel states
        if not incoming_phase and state != "Funnel":
            incoming_phase = project.get("project_otd_phase") or "Discovery"

        if str(incoming_otd or "").strip().lower().replace("-", " ") == "adjusted plan":
            if not adjusted:
                return False, None, "Adjusted End Date is required when OTD Status is Adjusted Plan"
            try:
                datetime.strptime(str(adjusted), "%Y-%m-%d")
            except ValueError:
                return False, None, "Adjusted End Date must be in YYYY-MM-DD format"
        else:
            # Preserve existing adjusted date — only let the Adjusted Plan form modify it
            adjusted = project.get("adjusted_end_date") or ""

        old_snapshot = project.copy()
        project["state"] = state
        project["project_status"] = incoming_status
        project["project_otd_status"] = incoming_otd
        project["project_otd_phase"] = incoming_phase or "Discovery"
        project["adjusted_end_date"] = adjusted
        project["status_comments"] = (incoming_comments or "").strip()
        project["updated_at"] = datetime.now().isoformat()

        for i, proj in enumerate(self.portfolio.get("projects", [])):
            if proj.get("project_id") == project_id:
                self.portfolio["projects"][i] = project
                break

        if not self._save_portfolio():
            return False, None, "Failed to save project"

        uid = user_id or os.getenv("CURRENT_USER", "Unknown")
        uname = os.getenv("CURRENT_USER_NAME", uid)
        for field in ["state", "project_status", "project_otd_status", "project_otd_phase", "adjusted_end_date", "status_comments"]:
            if old_snapshot.get(field) != project.get(field):
                self.audit_service.log_field_change(project_id, field, old_snapshot.get(field), project.get(field), uid, uname)

        return True, project, "Project status updated"

    def get_project_audit_trail(self, project_id: str) -> Dict[str, Any]:
        """
        Get audit trail for a project formatted for dashboard display.
        
        Args:
            project_id: Project ID
            
        Returns:
            Audit trail display data with entries and view_all flag
        """
        return self.audit_service.get_audit_trail_display(project_id, display_count=1)

    def _matches_search(self, project: Dict[str, Any], query: str) -> bool:
        if not query:
            return True
        q = self._norm(query)
        if not q:
            return True
        for f in self.SEARCH_FIELDS:
            value = self._norm(project.get(f, ""))
            if q in value:
                return True
        return False

    def _normalize_filter_compare_value(self, value: Any) -> str:
        if value is None:
            return ""
        s = str(value).strip()
        if not s:
            return ""
        if s.lower() == "(blank)":
            return ""
        return s

    def _resolve_filter_candidate(self, row: Dict[str, Any], key: str) -> Any:
        if key in ["project_status", "project_status_display"]:
            return self._normalize_project_status(row.get("project_status_display") or row.get("project_status"))
        if key in ["project_otd_status", "project_otd_status_display"]:
            return self._normalize_otd(row.get("project_otd_status_display") or row.get("project_otd_status"))
        return row.get(key, "")

    def _matches_filters(self, row: Dict[str, Any], filters: Dict[str, List[str]]) -> bool:
        for key, values in filters.items():
            if not values:
                continue

            candidate = self._normalize_filter_compare_value(self._resolve_filter_candidate(row, key))
            match = False

            for selected in values:
                selected_value = str(selected)
                if selected_value == self.BLANK_FILTER_TOKEN:
                    if candidate == "":
                        match = True
                        break
                elif candidate == self._normalize_filter_compare_value(selected_value):
                    match = True
                    break

            if not match:
                return False

        return True

    def _collect_filter_options(self, rows: List[Dict[str, Any]]) -> Dict[str, List[str]]:
        budget_fields = self.get_budget_field_names_for_ui()
        filter_fields = [
            "project_status_display",
            "initiative_name",
            "dt_team_accountable",
            "project_otd_status_display",
            "state",
            "project_otd_phase",
            "category",
            "business_unit",
            "primary_function",
            "region",
            "level_of_effort",
            "technology",
            budget_fields["current_budget"],
            budget_fields["next_budget"],
        ]

        options: Dict[str, List[str]] = {}
        for field in filter_fields:
            values = set()
            for row in rows:
                normalized = self._normalize_filter_compare_value(self._resolve_filter_candidate(row, field))
                values.add(normalized if normalized else self.BLANK_FILTER_TOKEN)

            sorted_values = sorted(
                values,
                key=lambda v: (v == self.BLANK_FILTER_TOKEN, str(v).lower()),
            )
            options[field] = sorted_values

        return options

    def _parse_sort_date(self, value: Any) -> Optional[datetime]:
        text = str(value or "").strip()
        if not text:
            return None
        try:
            return datetime.strptime(text[:10], "%Y-%m-%d")
        except Exception:
            return None

    def _parse_sort_number(self, value: Any) -> Optional[float]:
        text = str(value or "").strip()
        if not text:
            return None
        cleaned = text.replace(",", "").replace("$", "")
        try:
            return float(cleaned)
        except Exception:
            return None

    def _sort_value(self, row: Dict[str, Any], sort_key: str) -> Any:
        budget_fields = self.get_budget_field_names_for_ui()

        if sort_key == "end_date":
            return self._parse_sort_date(row.get("adjusted_end_date") or row.get("estd_end_date"))

        if sort_key == "budget":
            numeric_budget = self._parse_sort_number(row.get("budget"))
            return numeric_budget if numeric_budget is not None else str(row.get("budget") or "").strip().lower()

        if sort_key == "budget_plan":
            return str(row.get(budget_fields["current_budget"]) or "").strip().lower()

        if sort_key == "benefits_plan":
            return str(row.get(budget_fields["current_benefits"]) or "").strip().lower()

        if sort_key == "audit_trail":
            audit_payload = row.get("audit_trail") or {}
            if isinstance(audit_payload, dict):
                return int(audit_payload.get("total_count") or 0)
            return 0

        if sort_key == "actions":
            return ""

        value = row.get(sort_key)
        if sort_key in ["estd_end_date", "adjusted_end_date", "created_at", "updated_at"]:
            parsed = self._parse_sort_date(value)
            if parsed is not None:
                return parsed

        if sort_key in ["project_status_display", "project_status", "project_otd_status_display", "project_otd_status"]:
            return str(value or "").strip().lower()

        return str(value or "").strip().lower()

    def _sort_rows(self, rows: List[Dict[str, Any]], sort_key: str, sort_dir: str) -> List[Dict[str, Any]]:
        if not sort_key:
            return rows

        reverse = sort_dir == "desc"
        non_blank: List[Tuple[Any, Dict[str, Any]]] = []
        blank_rows: List[Dict[str, Any]] = []

        for row in rows:
            value = self._sort_value(row, sort_key)
            is_blank = value is None or (isinstance(value, str) and not value)
            if is_blank:
                blank_rows.append(row)
            else:
                non_blank.append((value, row))

        non_blank.sort(key=lambda item: item[0], reverse=reverse)
        return [row for _, row in non_blank] + blank_rows

    def get_search_suggestions(self, query: str, max_items: int = 6) -> List[str]:
        q = self._norm(query)
        if not q:
            return []
        candidates: List[str] = []
        for p in self.portfolio.get("projects", []):
            for f in self.SEARCH_FIELDS:
                value = str(p.get(f, "")).strip()
                if value:
                    candidates.append(value)
        uniq = sorted(set(candidates), key=lambda x: x.lower())
        starts = [x for x in uniq if x.lower().startswith(q)]
        contains = [x for x in uniq if q in x.lower() and x not in starts]
        return (starts + contains)[:max_items]

    def get_initiative_suggestions(self, query: str = "", max_items: int = 500) -> List[str]:
        """Return unique initiative names for elastic-search style initiative pickers."""
        q = self._norm(query)
        initiatives = {
            str(p.get("initiative_name", "")).strip()
            for p in self.portfolio.get("projects", [])
            if str(p.get("initiative_name", "")).strip()
        }
        ordered = sorted(initiatives, key=lambda x: x.lower())
        if q:
            starts = [x for x in ordered if x.lower().startswith(q)]
            contains = [x for x in ordered if q in x.lower() and x not in starts]
            ordered = starts + contains
        cap = max(1, int(max_items or 1))
        return ordered[:cap]

    def get_project_name_suggestions(self, query: str = "", max_items: int = 500) -> List[str]:
        """Return unique project names for elastic-search style project pickers."""
        q = self._norm(query)
        project_names = {
            str(p.get("project_name", "")).strip()
            for p in self.portfolio.get("projects", [])
            if str(p.get("project_name", "")).strip()
        }
        ordered = sorted(project_names, key=lambda x: x.lower())
        if q:
            starts = [x for x in ordered if x.lower().startswith(q)]
            contains = [x for x in ordered if q in x.lower() and x not in starts]
            ordered = starts + contains
        cap = max(1, int(max_items or 1))
        return ordered[:cap]

    def get_project_name_duplicate_matches(self, query: str, max_items: int = 3) -> List[Dict[str, Any]]:
        """Return likely duplicate project records for a project-name query."""
        raw_query = str(query or "").strip()
        normalized_query = self._norm(raw_query)
        if not normalized_query:
            return []

        query_tokens = [token for token in normalized_query.split() if token]
        candidates: List[Dict[str, Any]] = []

        for project in self.portfolio.get("projects", []):
            project_name = str(project.get("project_name") or "").strip()
            if not project_name:
                continue

            normalized_name = self._norm(project_name)
            score = 0

            if normalized_name == normalized_query:
                score = 100
            elif normalized_name.startswith(normalized_query) or normalized_query.startswith(normalized_name):
                score = 92
            elif normalized_query in normalized_name or normalized_name in normalized_query:
                score = 86
            else:
                ratio_score = int(SequenceMatcher(None, normalized_query, normalized_name).ratio() * 100)
                score = max(score, ratio_score)

            if query_tokens:
                name_tokens = set(normalized_name.split())
                overlap_count = len([token for token in query_tokens if token in name_tokens])
                if overlap_count:
                    score = max(score, 70 + min(20, overlap_count * 10))

            if score < 72:
                continue

            status = self._normalize_project_status(project.get("project_status"))
            if status in ["R", "Y", "A"]:
                risk_badge = "At Risk"
            elif status == "G":
                risk_badge = "On Track"
            elif status == "Hold":
                risk_badge = "Hold"
            else:
                risk_badge = "NA"

            candidates.append({
                "score": score,
                "project_id": project.get("project_id", ""),
                "project_name": project_name,
                "initiative_name": str(project.get("initiative_name") or "").strip(),
                "state": str(project.get("state") or "").strip(),
                "risk_badge": risk_badge,
            })

        candidates.sort(key=lambda item: (-item.get("score", 0), str(item.get("project_name", "")).lower()))
        cap = max(1, int(max_items or 1))
        return candidates[:cap]

    def _apply_display_rules(self, project: Dict[str, Any]) -> Dict[str, Any]:
        row = dict(project)
        row["project_status_display"] = self._display_project_status(project)
        row["project_otd_status_display"] = self._normalize_otd(self._display_otd_status(project))
        row["audit_trail"] = self.get_project_audit_trail(project.get("project_id"))
        return row

    def _build_chart_data(self, rows: List[Dict[str, Any]]) -> Dict[str, Any]:
        status_counts: Dict[str, int] = {}
        level_of_effort_counts: Dict[str, int] = {}
        category_counts: Dict[str, int] = {}
        state_counts: Dict[str, int] = {}
        deadline_counts: Dict[str, int] = {
            "Overdue": 0,
            "0-7 Days": 0,
            "8-14 Days": 0,
            "15-30 Days": 0,
            "30+ Days": 0,
            "Unknown": 0,
        }
        today = datetime.now().date()

        for row in rows:
            status = self._normalize_project_status(row.get("project_status_display"))
            status_counts[status] = status_counts.get(status, 0) + 1

            loe = str(row.get("level_of_effort") or "TBD").strip() or "TBD"
            level_of_effort_counts[loe] = level_of_effort_counts.get(loe, 0) + 1

            category = str(row.get("category") or "Unspecified").strip() or "Unspecified"
            category_counts[category] = category_counts.get(category, 0) + 1

            state = str(row.get("state") or "Unspecified").strip() or "Unspecified"
            state_counts[state] = state_counts.get(state, 0) + 1

            target_date = row.get("adjusted_end_date") or row.get("estd_end_date")
            target_str = str(target_date or "").strip()
            if not target_str:
                deadline_counts["Unknown"] += 1
            else:
                try:
                    deadline = datetime.strptime(target_str[:10], "%Y-%m-%d").date()
                    days_left = (deadline - today).days
                    if days_left < 0:
                        deadline_counts["Overdue"] += 1
                    elif days_left <= 7:
                        deadline_counts["0-7 Days"] += 1
                    elif days_left <= 14:
                        deadline_counts["8-14 Days"] += 1
                    elif days_left <= 30:
                        deadline_counts["15-30 Days"] += 1
                    else:
                        deadline_counts["30+ Days"] += 1
                except Exception:
                    deadline_counts["Unknown"] += 1

        preferred_status_order = ["G", "Y", "R", "Hold", "NA"]
        status_labels = [s for s in preferred_status_order if s in status_counts]
        status_labels.extend(sorted([s for s in status_counts.keys() if s not in status_labels], key=lambda x: x.lower()))

        preferred_loe_order = ["Small", "Medium", "Large", "XLarge", "Extra Large", "TBD"]
        loe_labels = [l for l in preferred_loe_order if l in level_of_effort_counts]
        loe_labels.extend(sorted([l for l in level_of_effort_counts.keys() if l not in loe_labels], key=lambda x: x.lower()))

        category_labels = sorted(category_counts.keys(), key=lambda x: x.lower())

        preferred_state_order = ["Active", "Evaluation", "Funnel", "Closure", "Hold", "Cancelled"]
        state_labels = [s for s in preferred_state_order if s in state_counts]
        state_labels.extend(sorted([s for s in state_counts.keys() if s not in state_labels], key=lambda x: x.lower()))

        return {
            "status": {
                "labels": status_labels,
                "values": [status_counts[s] for s in status_labels],
            },
            "level_of_effort": {
                "labels": loe_labels,
                "values": [level_of_effort_counts[l] for l in loe_labels],
            },
            "category": {
                "labels": category_labels,
                "values": [category_counts[c] for c in category_labels],
            },
            "state": {
                "labels": state_labels,
                "values": [state_counts[s] for s in state_labels],
            },
            "deadline_proximity": {
                "labels": ["Overdue", "0-7 Days", "8-14 Days", "15-30 Days", "30+ Days", "Unknown"],
                "values": [deadline_counts[label] for label in ["Overdue", "0-7 Days", "8-14 Days", "15-30 Days", "30+ Days", "Unknown"]],
            },
        }

    def _build_summary_cards(self, rows: List[Dict[str, Any]], chart_data: Dict[str, Any]) -> Dict[str, Any]:
        total_projects = len(rows)

        status_labels = chart_data.get("status", {}).get("labels", [])
        status_values = chart_data.get("status", {}).get("values", [])
        status_counts = {
            str(label): int(status_values[idx])
            for idx, label in enumerate(status_labels)
            if idx < len(status_values)
        }

        deadline_labels = chart_data.get("deadline_proximity", {}).get("labels", [])
        deadline_values = chart_data.get("deadline_proximity", {}).get("values", [])
        deadline_counts = {
            str(label): int(deadline_values[idx])
            for idx, label in enumerate(deadline_labels)
            if idx < len(deadline_values)
        }

        today = datetime.now().date()
        updated_this_month = 0
        for row in rows:
            updated_at = str(row.get("updated_at") or "").strip()
            if not updated_at:
                continue

            parsed_date = None
            try:
                parsed_date = datetime.fromisoformat(updated_at).date()
            except Exception:
                try:
                    parsed_date = datetime.strptime(updated_at[:10], "%Y-%m-%d").date()
                except Exception:
                    parsed_date = None

            if parsed_date and parsed_date.year == today.year and parsed_date.month == today.month:
                updated_this_month += 1

        update_compliance_pct = round((updated_this_month * 100 / total_projects), 1) if total_projects else 0.0

        return {
            "total_projects": total_projects,
            "on_track": status_counts.get("G", 0),
            "at_risk": status_counts.get("Y", 0),
            "off_track": status_counts.get("R", 0),
            "overdue": deadline_counts.get("Overdue", 0),
            "update_compliance": {
                "percent": update_compliance_pct,
                "updated_this_month": updated_this_month,
                "total_projects": total_projects,
            },
        }

    def get_dashboard_data(
        self,
        query: str = "",
        filters: Optional[Dict[str, List[str]]] = None,
        page: int = 1,
        page_size: int = 10,
        grouped: bool = True,
        sort_key: str = "",
        sort_dir: str = "asc",
    ) -> Dict[str, Any]:
        filters = filters or {}
        page_size = max(10, min(100, int(page_size)))
        if page_size % 10 != 0:
            page_size = 10

        base_rows = [
            self._apply_display_rules(p)
            for p in self.portfolio.get("projects", [])
            if self._matches_search(p, query)
        ]
        rows = [row for row in base_rows if self._matches_filters(row, filters)]
        filter_options = self._collect_filter_options(rows)
        chart_data = self._build_chart_data(rows)
        summary_cards = self._build_summary_cards(rows, chart_data)

        if grouped:
            groups_map: Dict[str, List[Dict[str, Any]]] = {}
            for row in rows:
                initiative = row.get("initiative_name") or "(No Initiative)"
                groups_map.setdefault(initiative, []).append(row)
            groups = [{"initiative_name": k, "project_count": len(v), "projects": v} for k, v in groups_map.items()]

            if sort_key:
                if sort_key == "initiative_name":
                    groups.sort(key=lambda g: (g.get("initiative_name") or "").lower(), reverse=sort_dir == "desc")
                else:
                    for group in groups:
                        group["projects"] = self._sort_rows(group.get("projects", []), sort_key, sort_dir)
                    groups.sort(key=lambda g: (g.get("initiative_name") or "").lower())
            else:
                groups.sort(key=lambda g: (g.get("initiative_name") or "").lower())

            total = len(groups)
            start = (max(1, page) - 1) * page_size
            items = groups[start:start + page_size]
            return {
                "grouped": True,
                "items": items,
                "total": total,
                "page": max(1, page),
                "page_size": page_size,
                "pages": max(1, (total + page_size - 1) // page_size),
                "project_total": len(rows),
                "chart_data": chart_data,
                "summary_cards": summary_cards,
                "filter_options": filter_options,
            }

        if sort_key:
            rows = self._sort_rows(rows, sort_key, sort_dir)
        else:
            rows.sort(key=lambda r: ((r.get("initiative_name") or "").lower(), (r.get("project_name") or "").lower()))
        total_rows = len(rows)
        start = (max(1, page) - 1) * page_size
        items = rows[start:start + page_size]
        return {
            "grouped": False,
            "items": items,
            "total": total_rows,
            "page": max(1, page),
            "page_size": page_size,
            "pages": max(1, (total_rows + page_size - 1) // page_size),
            "project_total": total_rows,
            "chart_data": chart_data,
            "summary_cards": summary_cards,
            "filter_options": filter_options,
        }

    def run_notification_cycle(self) -> Dict[str, Any]:
        notifications: List[Dict[str, Any]] = []
        today = datetime.now().date()
        month_start = today.replace(day=1)
        days_since_month_start = (today - month_start).days

        for p in self.portfolio.get("projects", []):
            owner = p.get("dt_team_accountable", "")
            target_date = p.get("adjusted_end_date") or p.get("estd_end_date")

            if days_since_month_start % 28 == 0:
                notifications.append({
                    "type": "monthly_status_reminder",
                    "channel": "email",
                    "recipient": owner,
                    "project_id": p.get("project_id"),
                    "project_name": p.get("project_name"),
                    "created_at": datetime.now().isoformat(),
                })

            if not target_date:
                continue

            try:
                deadline = datetime.strptime(str(target_date), "%Y-%m-%d").date()
            except ValueError:
                continue

            delta = (deadline - today).days
            if delta < 0:
                notifications.append({
                    "type": "breach_alert",
                    "channel": "email",
                    "recipient": owner,
                    "project_id": p.get("project_id"),
                    "project_name": p.get("project_name"),
                    "days_to_deadline": delta,
                    "created_at": datetime.now().isoformat(),
                })
            elif delta <= 28 and delta % 5 == 0:
                notifications.append({
                    "type": "proactive_deadline_warning",
                    "channel": "email",
                    "recipient": owner,
                    "project_id": p.get("project_id"),
                    "project_name": p.get("project_name"),
                    "days_to_deadline": delta,
                    "created_at": datetime.now().isoformat(),
                })

        existing: List[Dict[str, Any]] = []
        if self.notification_log_file.exists():
            try:
                loaded = json.loads(self.notification_log_file.read_text(encoding="utf-8"))
                if isinstance(loaded, list):
                    existing = loaded
            except Exception:
                existing = []

        existing.extend(notifications)
        self._write_json(self.notification_log_file, existing)
        return {"count": len(notifications), "notifications": notifications}
